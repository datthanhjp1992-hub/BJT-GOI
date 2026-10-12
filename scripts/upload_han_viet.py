"""
Ghi âm Hán Việt (Vocabulary.han_viet) lên Supabase thẳng từ máy local.

Nhanh hơn nhập qua màn SC07b trên server free: 1 câu SELECT + bulk_update theo
lô, không full_clean() từng dòng, không giới hạn 2.000 dòng/file.

File đầu vào: .xlsx (sheet đầu tiên) hoặc .csv, cần 3 cột word, reading,
han_viet — cột khác (id, meaning_vi, ghi_chu...) bỏ qua. Dùng được thẳng file
soát `Claude outputs/han_viet_soat.xlsx` hoặc các file han_viet_import_*.csv.

    # 0. cột han_viet phải có trên Supabase trước (migration 0004)
    python scripts/deploy_supabase.py --apply --yes

    # 1. chạy thử — so file với DB, KHÔNG ghi gì (mặc định)
    python scripts/upload_han_viet.py "Claude outputs/han_viet_soat.xlsx"

    # 2. ghi thật, người sửa là tài khoản dat
    python scripts/upload_han_viet.py "Claude outputs/han_viet_soat.xlsx" --apply --as-user dat

Khớp theo CẶP (word, reading) đúng như trên DB. File là nguồn chuẩn: ô
han_viet trống sẽ ghi trống. Chỉ ghi dòng có giá trị khác DB, nên chạy lại bao
nhiêu lần cũng ra cùng kết quả. Toàn bộ trong MỘT transaction.
"""
from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

from _vocab_common import DIM, GREEN, RED, YELLOW, act_as, say, setup_django

COLUMNS = ("word", "reading", "han_viet")


def read_rows(path: Path):
    """[(số dòng, word, reading, han_viet)] từ .csv / .xlsx."""
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as f:
            table = list(csv.reader(f))
    else:
        from openpyxl import load_workbook

        sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        table = [["" if v is None else str(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    if not table:
        return []
    headers = [h.strip() for h in table[0]]
    missing = [c for c in COLUMNS if c not in headers]
    if missing:
        say(RED, "LỖI", f"{path.name} thiếu cột: {', '.join(missing)}. Có: {', '.join(headers)}")
        sys.exit(1)
    idx = [headers.index(c) for c in COLUMNS]
    out = []
    for n, row in enumerate(table[1:], start=2):
        word, reading, hv = (row[i].strip() if i < len(row) else "" for i in idx)
        if word or reading:
            out.append((n, word, reading, hv))
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description="Ghi âm Hán Việt lên Supabase từ file .xlsx/.csv.")
    ap.add_argument("files", type=Path, nargs="+", help="File .xlsx / .csv có cột word, reading, han_viet.")
    ap.add_argument("--apply", action="store_true", help="Ghi thật. Không có cờ này thì chỉ chạy thử.")
    ap.add_argument("--as-user", metavar="USERNAME", help="Ghi updated_by là user này.")
    ap.add_argument("--settings", default="config.settings.supabase",
                    help="DJANGO_SETTINGS_MODULE (mặc định config.settings.supabase).")
    args = ap.parse_args()

    for path in args.files:
        if not path.exists():
            say(RED, "LỖI", f"Không thấy file {path}")
            sys.exit(1)

    setup_django(args.settings)
    from django.db import connection, transaction
    from django.utils import timezone

    from apps.core import middleware
    from apps.vocabulary.models import Vocabulary

    with connection.cursor() as c:
        cols = {col.name for col in connection.introspection.get_table_description(c, Vocabulary._meta.db_table)}
    if "han_viet" not in cols:
        say(RED, "LỖI", "DB chưa có cột han_viet — chạy trước: python scripts/deploy_supabase.py --apply --yes")
        sys.exit(1)

    say(YELLOW if args.apply else GREEN, "MODE", "GHI THẬT" if args.apply else "CHẠY THỬ (không ghi gì)")
    act_as(args.as_user)

    max_len = Vocabulary._meta.get_field("han_viet").max_length
    wanted, problems = {}, []
    for path in args.files:
        rows = read_rows(path)
        say(DIM, "FILE", f"{path} — {len(rows)} dòng")
        for n, word, reading, hv in rows:
            where = f"{path.name}:{n} {word}"
            if len(hv) > max_len:
                problems.append(f"{where}: han_viet dài {len(hv)} > {max_len} ký tự")
            elif (word, reading) in wanted and wanted[(word, reading)] != hv:
                problems.append(f"{where}: trùng từ với giá trị khác")
            else:
                wanted[(word, reading)] = hv
    if problems:
        for p in problems[:20]:
            say(RED, "LỖI", p)
        say(RED, "DỪNG", f"{len(problems)} dòng lỗi — không ghi gì.")
        sys.exit(1)

    by_pair = {(v.word, v.reading): v for v in Vocabulary.objects.only("id", "word", "reading", "han_viet")}
    not_found = [k for k in wanted if k not in by_pair]
    changed = []
    for key, hv in wanted.items():
        vocab = by_pair.get(key)
        if vocab is not None and vocab.han_viet != hv:
            vocab.han_viet = hv
            changed.append(vocab)

    for word, reading in not_found[:20]:
        say(YELLOW, "KHÔNG THẤY", f"{word}（{reading}）")
    say(DIM, "TỔNG", f"{len(wanted)} từ trong file · {len(changed)} cần ghi · "
        f"{len(wanted) - len(changed) - len(not_found)} đã đúng · {len(not_found)} không có trên DB")

    if not args.apply:
        say(GREEN, "XONG", "Chạy thử xong. Thêm --apply để ghi thật.")
        return
    now, user = timezone.now(), middleware._thread_locals.user
    for vocab in changed:
        vocab.updated_at = now  # bulk_update không gọi save() nên tự điền cột audit
        vocab.updated_by = user
    with transaction.atomic():
        Vocabulary.objects.bulk_update(changed, ["han_viet", "updated_at", "updated_by"], batch_size=500)
    say(GREEN, "XONG", f"Đã ghi {len(changed)} từ.")


if __name__ == "__main__":
    main()
