"""
Cập nhật nghĩa (meaning_vi) và âm Hán Việt / tiếng Anh (han_viet) của từ vựng
lên Supabase thẳng từ máy local, theo file .csv / .xlsx.

Dùng cho các file xuất từ màn Quản trị → Dữ liệu (cột word, reading,
meaning_vi, han_viet, audio_url...) đã sửa xong, vd
`Claude outputs/vocabulary_cap_nhat_phan1.csv`. Thay cho việc nhập qua web
khi file lớn: 1 câu SELECT + bulk_update theo lô, không giới hạn 2.000 dòng.

    # 1. chạy thử — so file với DB, KHÔNG ghi gì (mặc định)
    python scripts/upload_vocab_fields.py "Claude outputs/vocabulary_cap_nhat_phan1.csv"

    # 2. ghi thật, người sửa là tài khoản dat
    python scripts/upload_vocab_fields.py "Claude outputs/vocabulary_cap_nhat_phan1.csv" --apply --as-user dat

Khớp theo CẶP (word, reading) đúng như trên DB, chỉ với từ chưa bị xoá mềm.
Chỉ cập nhật các cột có trong file (meaning_vi và/hoặc han_viet); cột khác bỏ
qua. meaning_vi trống thì giữ nguyên DB (cột bắt buộc), han_viet trống thì ghi
trống. Chỉ ghi dòng có giá trị khác DB, nên chạy lại bao nhiêu lần cũng ra
cùng kết quả. Toàn bộ trong MỘT transaction.
"""
from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

from _vocab_common import DIM, GREEN, RED, YELLOW, act_as, say, setup_django

KEYS = ("word", "reading")
FIELDS = ("meaning_vi", "han_viet")


def read_table(path: Path):
    """(headers, [(số dòng, {cột: giá trị})]) từ .csv / .xlsx."""
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as f:
            table = list(csv.reader(f))
    else:
        from openpyxl import load_workbook

        sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        table = [["" if v is None else str(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    if not table:
        return [], []
    headers = [h.strip() for h in table[0]]
    out = []
    for n, row in enumerate(table[1:], start=2):
        data = {h: (row[i].strip() if i < len(row) else "") for i, h in enumerate(headers)}
        if data.get("word") or data.get("reading"):
            out.append((n, data))
    return headers, out


def main() -> None:
    ap = argparse.ArgumentParser(description="Cập nhật meaning_vi / han_viet lên Supabase từ file .xlsx/.csv.")
    ap.add_argument("files", type=Path, nargs="+", help="File .xlsx / .csv có cột word, reading và meaning_vi/han_viet.")
    ap.add_argument("--apply", action="store_true", help="Ghi thật. Không có cờ này thì chỉ chạy thử.")
    ap.add_argument("--as-user", metavar="USERNAME", help="Ghi updated_by là user này.")
    ap.add_argument("--settings", default="config.settings.supabase",
                    help="DJANGO_SETTINGS_MODULE (mặc định config.settings.supabase).")
    ap.add_argument("--show", type=int, default=10, metavar="N", help="In thử N thay đổi đầu tiên (mặc định 10).")
    args = ap.parse_args()

    for path in args.files:
        if not path.exists():
            say(RED, "LỖI", f"Không thấy file {path}")
            sys.exit(1)

    setup_django(args.settings)
    from django.db import transaction
    from django.utils import timezone

    from apps.core import middleware
    from apps.vocabulary.models import Vocabulary

    say(YELLOW if args.apply else GREEN, "MODE", "GHI THẬT" if args.apply else "CHẠY THỬ (không ghi gì)")
    act_as(args.as_user)

    max_len = {f: Vocabulary._meta.get_field(f).max_length for f in FIELDS}
    wanted, problems, used_fields = {}, [], set()
    for path in args.files:
        headers, rows = read_table(path)
        missing = [k for k in KEYS if k not in headers]
        fields = [f for f in FIELDS if f in headers]
        if missing or not fields:
            say(RED, "LỖI", f"{path.name} cần cột word, reading và ít nhất một trong {', '.join(FIELDS)}. "
                            f"Có: {', '.join(headers)}")
            sys.exit(1)
        used_fields.update(fields)
        say(DIM, "FILE", f"{path} — {len(rows)} dòng, cập nhật: {', '.join(fields)}")
        for n, data in rows:
            key = (data["word"], data["reading"])
            where = f"{path.name}:{n} {data['word']}"
            values = {}
            for f in fields:
                if f == "meaning_vi" and not data[f]:
                    continue  # cột bắt buộc: ô trống = giữ nguyên
                if len(data[f]) > max_len[f]:
                    problems.append(f"{where}: {f} dài {len(data[f])} > {max_len[f]} ký tự")
                values[f] = data[f]
            old = wanted.get(key)
            if old is not None and any(old.get(f, v) != v for f, v in values.items()):
                problems.append(f"{where}: trùng từ với giá trị khác")
            wanted.setdefault(key, {}).update(values)
    if problems:
        for p in problems[:20]:
            say(RED, "LỖI", p)
        say(RED, "DỪNG", f"{len(problems)} dòng lỗi — không ghi gì.")
        sys.exit(1)

    by_pair = {(v.word, v.reading): v for v in Vocabulary.objects.only("id", "word", "reading", *FIELDS)}
    not_found = [k for k in wanted if k not in by_pair]
    changed, counts, shown = [], dict.fromkeys(FIELDS, 0), 0
    for key, values in wanted.items():
        vocab = by_pair.get(key)
        if vocab is None:
            continue
        diff = {f: v for f, v in values.items() if getattr(vocab, f) != v}
        if not diff:
            continue
        for f, v in diff.items():
            if shown < args.show:
                say(DIM, "ĐỔI", f"{vocab.word}（{vocab.reading}）{f}: {getattr(vocab, f)!r} → {v!r}")
                shown += 1
            setattr(vocab, f, v)
            counts[f] += 1
        changed.append(vocab)

    for word, reading in not_found[:20]:
        say(YELLOW, "KHÔNG THẤY", f"{word}（{reading}）")
    say(DIM, "TỔNG", f"{len(wanted)} từ trong file · {len(changed)} cần ghi "
        f"({', '.join(f'{f}: {counts[f]}' for f in FIELDS if f in used_fields)}) · "
        f"{len(wanted) - len(changed) - len(not_found)} đã đúng · {len(not_found)} không có trên DB")

    if not args.apply:
        say(GREEN, "XONG", "Chạy thử xong. Thêm --apply để ghi thật.")
        return
    if not changed:
        say(GREEN, "XONG", "Không có gì cần ghi.")
        return
    now, user = timezone.now(), middleware._thread_locals.user
    for vocab in changed:
        vocab.updated_at = now  # bulk_update không gọi save() nên tự điền cột audit
        vocab.updated_by = user
    update_fields = [f for f in FIELDS if counts[f]] + ["updated_at", "updated_by"]
    with transaction.atomic():
        Vocabulary.objects.bulk_update(changed, update_fields, batch_size=500)
    say(GREEN, "XONG", f"Đã ghi {len(changed)} từ.")


if __name__ == "__main__":
    main()
