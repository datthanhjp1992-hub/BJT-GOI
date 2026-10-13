"""
Sửa từ vựng ĐÃ CÓ trên Supabase thẳng từ máy local: nghĩa, âm Hán Việt, audio.

Nhanh hơn nhập qua màn SC07b trên server free: 1 câu SELECT + bulk_update theo
lô, không full_clean() từng dòng, không giới hạn 2.000 dòng/file. Không tạo từ
mới (từ mới dùng scripts/upload_vocab.py).

File đầu vào: .xlsx (sheet đầu tiên) hoặc .csv. Bắt buộc cột word, reading và
ít nhất một trong meaning_vi, han_viet, audio_url — chỉ cột CÓ trong file mới
được ghi, cột khác (id, ghi_chu...) bỏ qua.

    # 1. chạy thử — so file với DB, KHÔNG ghi gì (mặc định)
    .venv\\Scripts\\python.exe scripts/update_vocab.py vocabulary_cap_nhat.csv

    # 2. ghi thật, người sửa là tài khoản dat
    .venv\\Scripts\\python.exe scripts/update_vocab.py vocabulary_cap_nhat.csv --apply --as-user dat

Khớp theo CẶP (word, reading) đúng như trên DB (từ đã xoá mềm không tính). File
là nguồn chuẩn: ô trống sẽ ghi trống (riêng meaning_vi bắt buộc có). Chỉ ghi
dòng có giá trị khác DB, nên chạy lại bao nhiêu lần cũng ra cùng kết quả.
Toàn bộ trong MỘT transaction.
"""
from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

from _vocab_common import DIM, GREEN, RED, YELLOW, act_as, say, setup_django

KEYS = ("word", "reading")
FIELDS = ("meaning_vi", "han_viet", "audio_url")


def read_table(path: Path):
    """(headers, các dòng) từ .csv / .xlsx."""
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as f:
            table = list(csv.reader(f))
    else:
        from openpyxl import load_workbook

        sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        table = [["" if v is None else str(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    if not table:
        return [], []
    return [h.strip() for h in table[0]], table[1:]


def main() -> None:
    ap = argparse.ArgumentParser(description="Sửa nghĩa / âm Hán Việt / audio của từ vựng đã có trên Supabase.")
    ap.add_argument("files", type=Path, nargs="+", help="File .xlsx / .csv có cột word, reading + cột cần sửa.")
    ap.add_argument("--apply", action="store_true", help="Ghi thật. Không có cờ này thì chỉ chạy thử.")
    ap.add_argument("--as-user", metavar="USERNAME", help="Ghi updated_by là user này.")
    ap.add_argument("--settings", default="config.settings.supabase",
                    help="DJANGO_SETTINGS_MODULE (mặc định config.settings.supabase).")
    args = ap.parse_args()

    for path in args.files:
        if not path.exists():
            say(RED, "LỖI", f"Không thấy file {path}")
            sys.exit(1)

    setup_django(args.settings)
    from django.core.exceptions import ValidationError
    from django.core.validators import URLValidator
    from django.db import transaction
    from django.utils import timezone

    from apps.core import middleware
    from apps.vocabulary.models import Vocabulary

    say(YELLOW if args.apply else GREEN, "MODE", "GHI THẬT" if args.apply else "CHẠY THỬ (không ghi gì)")
    act_as(args.as_user)

    max_len = {f: Vocabulary._meta.get_field(f).max_length for f in FIELDS}
    check_url = URLValidator()
    wanted, problems, fields = {}, [], set()
    for path in args.files:
        headers, rows = read_table(path)
        missing = [c for c in KEYS if c not in headers]
        cols = [f for f in FIELDS if f in headers]
        if missing or not cols:
            say(RED, "LỖI", f"{path.name} cần cột word, reading và ít nhất một trong {', '.join(FIELDS)}. "
                            f"Có: {', '.join(headers)}")
            sys.exit(1)
        fields.update(cols)
        say(DIM, "FILE", f"{path} — {len(rows)} dòng, sửa cột: {', '.join(cols)}")
        for n, row in enumerate(rows, start=2):
            get = lambda c: (row[headers.index(c)] if headers.index(c) < len(row) else "").strip()
            word, reading = get("word"), get("reading")
            if not (word or reading):
                continue
            data = {c: get(c) for c in cols}
            where = f"{path.name}:{n} {word}"
            for c, value in data.items():
                if len(value) > max_len[c]:
                    problems.append(f"{where}: {c} dài {len(value)} > {max_len[c]} ký tự")
            if "meaning_vi" in data and not data["meaning_vi"]:
                problems.append(f"{where}: meaning_vi trống")
            if data.get("audio_url"):
                try:
                    check_url(data["audio_url"])
                except ValidationError:
                    problems.append(f"{where}: audio_url không phải URL")
            old = wanted.setdefault((word, reading), {})
            if any(old.get(c, v) != v for c, v in data.items()):
                problems.append(f"{where}: trùng từ với giá trị khác")
            old.update(data)
    if problems:
        for p in problems[:20]:
            say(RED, "LỖI", p)
        say(RED, "DỪNG", f"{len(problems)} dòng lỗi — không ghi gì.")
        sys.exit(1)

    by_pair = {(v.word, v.reading): v for v in Vocabulary.objects.only("id", "word", "reading", *FIELDS)}
    not_found = [k for k in wanted if k not in by_pair]
    changed, per_field = [], {f: 0 for f in sorted(fields)}
    for key, data in wanted.items():
        vocab = by_pair.get(key)
        if vocab is None:
            continue
        diff = [c for c, v in data.items() if (getattr(vocab, c) or "") != v]
        for c in diff:
            setattr(vocab, c, data[c])
            per_field[c] += 1
        if diff:
            changed.append(vocab)

    for word, reading in not_found[:20]:
        say(YELLOW, "KHÔNG THẤY", f"{word}（{reading}）")
    say(DIM, "TỔNG", f"{len(wanted)} từ trong file · {len(changed)} cần ghi · "
        f"{len(wanted) - len(changed) - len(not_found)} đã đúng · {len(not_found)} không có trên DB")
    say(DIM, "THEO CỘT", " · ".join(f"{c}: {n} đổi" for c, n in per_field.items()))

    if not args.apply:
        say(GREEN, "XONG", "Chạy thử xong. Thêm --apply để ghi thật.")
        return
    now, user = timezone.now(), middleware._thread_locals.user
    for vocab in changed:
        vocab.updated_at = now  # bulk_update không gọi save() nên tự điền cột audit
        vocab.updated_by = user
    with transaction.atomic():
        Vocabulary.objects.bulk_update(changed, sorted(fields) + ["updated_at", "updated_by"], batch_size=500)
    say(GREEN, "XONG", f"Đã ghi {len(changed)} từ.")


if __name__ == "__main__":
    main()
