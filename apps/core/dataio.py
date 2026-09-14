"""
Engine nhập/xuất dữ liệu bằng file cho khu quản trị SC07b.

Ý tưởng: KHÔNG viết tay 18 bộ form/parse. Cột của mỗi bảng được suy ra từ
`Model._meta` nên thêm/bớt field trong models.py là file mẫu, file xuất và
trình nhập tự đổi theo — không có nơi thứ hai phải sửa.

Quy ước cột:
- Bỏ hết khoá chính và 4 cột audit (created_by/created_at/updated_by/updated_at).
  Audit do `AuditableModel.save()` + `CurrentUserMiddleware` tự điền.
- Field thường  -> 1 cột trùng tên field.
- Khoá ngoại    -> dò theo KHOÁ TỰ NHIÊN của bảng đích, ĐÃ TRẢI PHẲNG qua nhiều
  cấp FK. Cột đặt tên `<field>__<đường dẫn>`, vd `vocabulary__word`,
  `wordlist__user__username`, `sheet__pdf_file`. Bảng đích không có khoá tự
  nhiên nào thì mới lùi về `<field>_id` (id thô).
- ManyToMany KHÔNG có cột. Cả 3 quan hệ M2M của repo đều đi qua through model
  (VocabularyTopic / UserWordlistWord / PracticeSheetWord) và 3 bảng nối đó tự
  nằm trong danh sách nhập được — nhập bảng nối là cách duy nhất giữ đủ audit.

BA CHẾ ĐỘ GHI (14/09/2026) — trước đây chỉ có chế độ đầu tiên:
- `insert`  Chỉ thêm mới. Dòng đã có bị BỎ QUA.
- `upsert`  Thêm + cập nhật. Dòng đã có được ghi đè theo giá trị trong file.
- `update`  Chỉ cập nhật. Dòng chưa có bị BỎ QUA (đếm riêng, không phải lỗi).
Không có chế độ xoá, và sẽ không thêm: xoá một từ vựng kéo theo ExampleSentence,
tiến độ học và wordlist của người dùng.

Riêng chế độ `update` chỉ đòi các CỘT KHOÁ; cột nào không có trong file thì
field đó giữ nguyên giá trị cũ. Nhờ vậy muốn sửa mỗi `meaning_vi` thì chỉ cần
file 3 cột `word,reading,meaning_vi`. Ngược lại, cột CÓ trong file mà ô để
trống nghĩa là "ghi giá trị rỗng/mặc định" — muốn giữ nguyên thì xoá hẳn cột đó.

MẪU GỘP "TỪ VỰNG ĐẦY ĐỦ" (14/09/2026): một từ vựng trải trên 3 bảng
(Vocabulary + VocabularyTopic + ExampleSentence) nên nhập theo từng bảng phải
tải lên 3 file đúng thứ tự. `FullVocabularyDataset` gộp lại thành MỘT file:
1 dòng = 1 từ, kèm cột `topics` và 3 cặp câu ví dụ. Xem cuối file.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field as dataclass_field

from django.apps import apps as django_apps
from django.core.exceptions import ValidationError
from django.db import models

# 6 app nội bộ = 18 bảng. django.contrib.auth KHÔNG đụng vào (quyết định của Dat,
# xem claude/db-schema-django.md).
LOCAL_APP_LABELS = (
    "core",
    "accounts",
    "vocabulary",
    "learning",
    "practice_sheets",
    "gamification",
)

AUDIT_FIELDS = ("created_by", "created_at", "updated_by", "updated_at")

# Tên sheet hướng dẫn trong file mẫu .xlsx — trùng với chuỗi đã dùng ở
# message.properties (practice_sheets.upload.error.missing_required_column).
GUIDE_SHEET_NAME = "Huong_dan"

# Chặn trên số dòng mỗi lần nhập. Mỗi dòng chạy full_clean() (có validate_unique)
# nên file quá to sẽ treo request; chia nhỏ file thay vì nâng số này.
MAX_IMPORT_ROWS = 2000

CSV_ENCODINGS = ("utf-8-sig", "cp932", "utf-8")

BOOL_TRUE = {"1", "true", "t", "yes", "y", "x", "co", "có", "dung", "đúng"}
BOOL_FALSE = {"0", "false", "f", "no", "n", "", "khong", "không", "sai"}

# ---------------------------------------------------------------------------
# Chế độ ghi
# ---------------------------------------------------------------------------

MODE_INSERT = "insert"
MODE_UPSERT = "upsert"
MODE_UPDATE = "update"
WRITE_MODES = (MODE_INSERT, MODE_UPSERT, MODE_UPDATE)
DEFAULT_MODE = MODE_INSERT

# Chế độ nào cần quyền gì. Không cho staff chỉ có quyền `add` chạy chế độ ghi đè.
MODE_PERMISSIONS = {
    MODE_INSERT: ("add",),
    MODE_UPSERT: ("add", "change"),
    MODE_UPDATE: ("change",),
}


def normalize_mode(value):
    """Chế độ lạ (người dùng sửa tay form) lùi về chế độ an toàn nhất."""
    value = (value or "").strip().lower()
    return value if value in WRITE_MODES else DEFAULT_MODE


def mode_writes_new_rows(mode):
    return mode in (MODE_INSERT, MODE_UPSERT)


def mode_writes_existing_rows(mode):
    return mode in (MODE_UPSERT, MODE_UPDATE)


# ---------------------------------------------------------------------------
# Khoá tự nhiên
# ---------------------------------------------------------------------------
# Dùng cho (1) cột tham chiếu FK, (2) phát hiện trùng, (3) tìm dòng cần cập nhật.
# Khai được cả field FK: `natural_key_paths()` tự trải phẳng qua bảng đích.

NATURAL_KEYS = {
    "accounts.user": ("username",),
    "core.mastercode": ("code_type", "code"),
    "vocabulary.topic": ("slug",),
    "vocabulary.vocabulary": ("word", "reading"),
    # Một từ không thể có hai câu ví dụ giống hệt nhau -> khoá là (từ, câu Nhật).
    # Thiếu khoá này thì nhập lại cùng file là nhân đôi toàn bộ câu ví dụ.
    "vocabulary.examplesentence": ("vocabulary", "sentence_jp"),
    "vocabulary.vocabularytopic": ("vocabulary", "topic"),
    # Một người không đặt hai danh sách trùng tên.
    "learning.userwordlist": ("user", "name"),
    "learning.userwordlistword": ("wordlist", "vocabulary"),
    "learning.uservocabularyprogress": ("user", "vocabulary"),
    # Một người không bắt đầu hai phiên học ở đúng cùng một mốc thời gian.
    "learning.studysession": ("user", "started_at"),
    # PracticeSheet là bản ghi của MỘT file PDF đã sinh ra -> đường dẫn file là
    # thứ duy nhất phân biệt được hai bản ghi của cùng một người.
    "practice_sheets.practicesheet": ("user", "pdf_file"),
    "practice_sheets.practicesheetword": ("sheet", "vocabulary"),
    "gamification.pointrule": ("action_code",),
    # Khoá của BadgeCategory là code_type (1 nhóm danh hiệu = 1 code_type bên
    # MasterCode), KHÔNG phải "code" — model này không có field tên code.
    "gamification.badgecategory": ("code_type",),
    "gamification.badgetier": ("category", "code"),
    "gamification.userpinnedbadge": ("user", "category"),
}

# Hai bảng NHẬT KÝ cố ý không có khoá tự nhiên, và đừng bịa ra một cái.
#   - Contribution: cùng một người gửi hai góp ý nội dung y hệt nhau vẫn là hai
#     góp ý khác nhau (gửi lại sau khi bị từ chối).
#   - UserPointTransaction: cùng một người được cộng cùng số điểm cho cùng hành
#     động nhiều lần là chuyện bình thường.
# Bịa khoá ở đây sẽ khiến chế độ "thêm + cập nhật" GHI ĐÈ một bản ghi lịch sử
# hợp lệ. Thay vào đó: hai bảng này chỉ cho nhập ở chế độ THÊM MỚI, và màn nhập
# cảnh báo rõ là nhập hai lần sẽ ra dữ liệu đôi.
TABLES_WITHOUT_NATURAL_KEY = frozenset(
    {
        "gamification.contribution",
        "gamification.userpointtransaction",
    }
)

# Chặn đệ quy khi trải phẳng khoá tự nhiên qua FK (A khoá theo B, B khoá theo A).
MAX_NATURAL_KEY_DEPTH = 3

# Cột có trong file nhưng để TRỐNG thì giữ nguyên giá trị cũ thay vì ghi rỗng.
# Chỉ dùng cho thứ không bao giờ được phép vô tình xoá mất.
BLANK_MEANS_UNCHANGED = {
    "accounts.user": ("password",),
}


class DataFileError(Exception):
    """Lỗi ở tầng file (sai định dạng, thiếu cột, quá số dòng)."""

    def __init__(self, message_key, **params):
        super().__init__(message_key)
        self.message_key = message_key
        self.params = params


# ---------------------------------------------------------------------------
# Danh sách bảng
# ---------------------------------------------------------------------------

def model_label(model):
    return f"{model._meta.app_label}.{model._meta.model_name}"


def _all_local_models():
    found = []
    for app_label in LOCAL_APP_LABELS:
        try:
            config = django_apps.get_app_config(app_label)
        except LookupError:  # pragma: no cover - app bị gỡ khỏi INSTALLED_APPS
            continue
        for model in config.get_models():
            if model._meta.auto_created or model._meta.proxy:
                continue
            found.append(model)
    return found


def importable_models():
    """18 bảng, xếp theo thứ tự phụ thuộc: bảng bị tham chiếu đứng trước.

    Nhập theo đúng thứ tự này thì khoá ngoại luôn dò ra. Vòng lặp tự tham chiếu
    (User.created_by trỏ về chính User) không tính là phụ thuộc.
    """
    models_list = _all_local_models()
    by_label = {model_label(m): m for m in models_list}
    remaining = dict(by_label)
    ordered = []
    while remaining:
        ready = []
        for label, model in remaining.items():
            deps = set()
            for f in model._meta.fields:
                if f.is_relation and f.related_model is not None:
                    dep = model_label(f.related_model)
                    if dep != label and dep in remaining:
                        deps.add(dep)
            if not deps:
                ready.append(label)
        if not ready:  # còn phụ thuộc vòng -> xả phần còn lại theo tên
            ready = sorted(remaining)
        for label in sorted(ready):
            ordered.append(remaining.pop(label))
    return ordered


def get_importable_model(label):
    """Tra model theo nhãn 'app_label.modelname'. None nếu không nằm trong 18 bảng."""
    label = (label or "").strip().lower()
    for model in _all_local_models():
        if model_label(model) == label:
            return model
    return None


def verbose_name(model):
    return str(model._meta.verbose_name).strip() or model._meta.model_name


# ---------------------------------------------------------------------------
# Cột
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class Column:
    header: str          # tên cột trong file
    field_name: str      # tên field trên model (mẫu gộp: tên khoá logic)
    lookup: str | None   # FK: đường dẫn dò trên bảng đích; None = id thô / field thường
    is_fk: bool
    required: bool
    hint: str            # mô tả ngắn để in vào sheet Huong_dan
    is_key: bool = False  # thuộc khoá tự nhiên -> bắt buộc ở chế độ "chỉ cập nhật"

    @property
    def label(self):
        return self.header


def _fields_exist(model, names):
    for name in names:
        try:
            model._meta.get_field(name)
        except Exception:
            return False
    return True


def natural_key_fields(model):
    """Bộ field (tên field trên chính model, có thể là FK) tạo nên khoá tự nhiên.

    Ưu tiên khoá đã khai ở NATURAL_KEYS; sau đó tới UniqueConstraint,
    unique_together, rồi field unique=True đơn lẻ. Rỗng = bảng không có khoá
    nào -> mọi dòng đều tính là mới, và không cập nhật được.

    NATURAL_KEYS gõ sai tên field thì BỎ QUA chứ không dùng: nếu tin nó, câu
    filter() lúc dò trùng sẽ nổ FieldError giữa lúc admin đang nhập.
    """
    label = model_label(model)
    if label in TABLES_WITHOUT_NATURAL_KEY:
        return ()
    keys = NATURAL_KEYS.get(label)
    if keys and _fields_exist(model, keys):
        return tuple(keys)
    for constraint in model._meta.constraints:
        fields = getattr(constraint, "fields", None)
        if fields:
            return tuple(fields)
    for together in model._meta.unique_together:
        return tuple(together)
    for f in model._meta.fields:
        if f.unique and not f.primary_key and f.editable:
            return (f.name,)
    return ()


# Tên cũ, giữ lại vì view/template/test đang gọi.
def duplicate_fields(model):
    """Bộ field dùng để nhận ra 'dòng này đã có trong DB'. Xem natural_key_fields()."""
    return natural_key_fields(model)


def natural_key_paths(model, _depth=0):
    """Khoá tự nhiên đã TRẢI PHẲNG thành đường dẫn lookup dùng được làm cột.

    UserWordlist khoá theo (user, name) mà user lại là FK -> trả về
    ("user__username", "name"). Nhờ vậy cột trỏ tới UserWordlist là
    `wordlist__user__username` + `wordlist__name` thay vì `wordlist_id` mà
    không ai điền nổi.

    Trả về () khi có mắt xích không trải được (bảng trung gian không có khoá tự
    nhiên) hoặc khi đi quá sâu — lúc đó bên gọi lùi về cột id thô.
    """
    if _depth >= MAX_NATURAL_KEY_DEPTH:
        return ()
    paths = []
    for name in natural_key_fields(model):
        f = model._meta.get_field(name)
        if f.is_relation:
            if f.related_model is None:  # pragma: no cover
                return ()
            sub = natural_key_paths(f.related_model, _depth + 1)
            if not sub:
                return ()
            paths.extend(f"{name}__{s}" for s in sub)
        else:
            paths.append(name)
    return tuple(paths)


def _field_hint(f):
    parts = [f.get_internal_type()]
    if getattr(f, "max_length", None):
        parts.append(f"tối đa {f.max_length}")
    if f.choices:
        parts.append("giá trị nằm trong MasterCode")
    if f.has_default():
        parts.append("bỏ trống = mặc định")
    elif f.null or f.blank:
        parts.append("được để trống")
    else:
        parts.append("bắt buộc")
    if f.help_text:
        parts.append(str(f.help_text))
    return " · ".join(parts)


def columns_for(model):
    """Danh sách cột của 1 bảng, đúng thứ tự khai trong models.py."""
    key_fields = set(natural_key_fields(model))
    columns = []
    for f in model._meta.fields:
        if f.primary_key or not f.editable or f.auto_created:
            continue
        if f.name in AUDIT_FIELDS:
            continue
        required = not (f.blank or f.null or f.has_default())
        is_key = f.name in key_fields
        if isinstance(f, models.ForeignKey):
            target_paths = natural_key_paths(f.related_model)
            if target_paths:
                for path in target_paths:
                    columns.append(
                        Column(
                            header=f"{f.name}__{path}",
                            field_name=f.name,
                            lookup=path,
                            is_fk=True,
                            required=required,
                            is_key=is_key,
                            hint=f"tham chiếu {verbose_name(f.related_model)} theo {path}",
                        )
                    )
            else:
                columns.append(
                    Column(
                        header=f"{f.name}_id",
                        field_name=f.name,
                        lookup=None,
                        is_fk=True,
                        required=required,
                        is_key=is_key,
                        hint=(
                            f"id của {verbose_name(f.related_model)} — bảng đó không có "
                            f"khoá tự nhiên nên phải điền id thô"
                        ),
                    )
                )
        else:
            columns.append(
                Column(
                    header=f.name,
                    field_name=f.name,
                    lookup=None,
                    is_fk=False,
                    required=required,
                    is_key=is_key,
                    hint=_field_hint(f),
                )
            )
    return columns


def headers_for(model):
    return [c.header for c in columns_for(model)]


def skipped_m2m_names(model):
    """M2M bị bỏ khỏi file — UI in ra để admin biết phải nhập bảng nối nào."""
    out = []
    for f in model._meta.many_to_many:
        through = getattr(f.remote_field, "through", None)
        if through is not None and not through._meta.auto_created:
            out.append((f.name, model_label(through)))
        else:  # pragma: no cover - repo không còn M2M ẩn nào
            out.append((f.name, ""))
    return out


# ---------------------------------------------------------------------------
# Đọc file
# ---------------------------------------------------------------------------

def _decode_csv(raw):
    for encoding in CSV_ENCODINGS:
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DataFileError("admin.data.error.encoding")


def _read_csv(raw):
    text = _decode_csv(raw)
    reader = csv.reader(io.StringIO(text, newline=""))
    rows = [list(r) for r in reader]
    if not rows:
        raise DataFileError("admin.data.error.empty_file")
    headers = [str(h).strip() for h in rows[0]]
    return headers, rows[1:]


def _read_xlsx(raw):
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover - openpyxl nằm trong requirements
        raise DataFileError("admin.data.error.xlsx_unavailable")
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheet = None
    for candidate in workbook.worksheets:
        if candidate.title != GUIDE_SHEET_NAME:
            sheet = candidate
            break
    if sheet is None:
        raise DataFileError("admin.data.error.empty_file")
    rows = []
    for row in sheet.iter_rows(values_only=True):
        rows.append(["" if v is None else v for v in row])
    workbook.close()
    if not rows:
        raise DataFileError("admin.data.error.empty_file")
    headers = [str(h).strip() for h in rows[0]]
    return headers, rows[1:]


def read_table(raw, filename):
    """(headers, rows) từ nội dung file. Chấp nhận .csv / .xlsx / .xlsm."""
    name = (filename or "").lower()
    if name.endswith(".csv"):
        headers, rows = _read_csv(raw)
    elif name.endswith(".xlsx") or name.endswith(".xlsm"):
        headers, rows = _read_xlsx(raw)
    else:
        raise DataFileError("admin.data.error.unsupported_format")
    # Bỏ dòng trống hoàn toàn ở cuối file (Excel hay để lại)
    while rows and all(str(v).strip() == "" for v in rows[-1]):
        rows.pop()
    if not rows:
        raise DataFileError("admin.data.error.empty_file")
    if len(rows) > MAX_IMPORT_ROWS:
        raise DataFileError("admin.data.error.too_many_rows", max_rows=MAX_IMPORT_ROWS, row_count=len(rows))
    return headers, rows


# ---------------------------------------------------------------------------
# Ghi file
# ---------------------------------------------------------------------------

def write_csv(headers, rows):
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    # BOM để Excel bản tiếng Nhật mở UTF-8 không ra mojibake.
    return buffer.getvalue().encode("utf-8-sig")


def write_xlsx(headers, rows, guide_rows=None, sheet_title="Data"):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = (sheet_title or "Data")[:31]
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    for row in rows:
        sheet.append(["" if v is None else v for v in row])
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = max(
            12, min(40, len(str(header)) + 6)
        )
    if guide_rows:
        guide = workbook.create_sheet(GUIDE_SHEET_NAME)
        for row in guide_rows:
            guide.append(list(row))
        for cell in guide[1]:
            cell.font = Font(bold=True)
        guide.column_dimensions["A"].width = 32
        guide.column_dimensions["B"].width = 14
        guide.column_dimensions["C"].width = 70
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def export_rows(model, queryset=None):
    """Dữ liệu hiện có của 1 bảng, đúng bộ cột của file mẫu."""
    columns = columns_for(model)
    queryset = model._default_manager.all() if queryset is None else queryset
    fk_names = {c.field_name for c in columns if c.is_fk}
    if fk_names:
        queryset = queryset.select_related(*sorted(fk_names))
    rows = []
    for obj in queryset.iterator():
        row = []
        for column in columns:
            if column.is_fk:
                related = getattr(obj, column.field_name, None)
                if related is None:
                    row.append("")
                elif column.lookup:
                    row.append(_follow_path(related, column.lookup))
                else:
                    row.append(related.pk)
            else:
                value = getattr(obj, column.field_name, "")
                row.append("" if value is None else value)
        rows.append(row)
    return rows


def _follow_path(obj, path):
    """`_follow_path(link.wordlist, "user__username")` -> "nguoihoc"."""
    current = obj
    for part in path.split("__"):
        if current is None:
            return ""
        current = getattr(current, part, None)
    return "" if current is None else current


def _guide_mode_rows():
    return [
        [],
        ["Chế độ ghi", "", "Chọn ngay trên màn nhập, không nằm trong file này."],
        ["Chỉ thêm mới", "", "Dòng đã có trong hệ thống bị bỏ qua, dữ liệu cũ giữ nguyên."],
        ["Thêm + cập nhật", "", "Dòng đã có được ghi đè theo giá trị trong file."],
        [
            "Chỉ cập nhật",
            "",
            "Chỉ cần các cột KHOÁ; cột nào không có trong file thì field đó giữ "
            "nguyên. Cột CÓ trong file mà ô để trống = ghi giá trị rỗng/mặc định.",
        ],
    ]


def guide_rows(model):
    out = [["Cột", "Bắt buộc", "Ghi chú"]]
    key_headers = [c.header for c in columns_for(model) if c.is_key]
    for column in columns_for(model):
        note = column.hint
        if column.is_key:
            note = f"[KHOÁ] {note}"
        out.append([column.header, "x" if column.required else "", note])
    for name, through in skipped_m2m_names(model):
        out.append([name, "", f"KHÔNG nhập ở đây — nhập bảng nối {through}"])
    out.extend(_guide_mode_rows())
    out.append([])
    if key_headers:
        out.append(["Khoá chống trùng", "", ", ".join(key_headers)])
    else:
        out.append(
            [
                "Khoá chống trùng",
                "",
                "KHÔNG CÓ — mọi dòng đều tính là mới, nhập hai lần sẽ ra dữ liệu "
                "đôi, và bảng này không cập nhật được bằng file.",
            ]
        )
    return out


# ---------------------------------------------------------------------------
# Phân tích file trước khi ghi
# ---------------------------------------------------------------------------

STATUS_NEW = "new"
STATUS_UPDATE = "update"
STATUS_DUPLICATE = "duplicate"
STATUS_MISSING = "missing"
STATUS_ERROR = "error"


@dataclass
class RowResult:
    number: int                      # số dòng thật trong file (header là dòng 1)
    status: str
    display: str = ""
    errors: list = dataclass_field(default_factory=list)
    data: dict = dataclass_field(default_factory=dict)
    pk: object = None                # bản ghi sẽ bị ghi đè (chế độ cập nhật)


@dataclass
class ImportReport:
    model_label: str
    model_name: str
    mode: str = DEFAULT_MODE
    headers: list = dataclass_field(default_factory=list)
    missing_columns: list = dataclass_field(default_factory=list)
    unknown_columns: list = dataclass_field(default_factory=list)
    ignored_columns: list = dataclass_field(default_factory=list)
    rows: list = dataclass_field(default_factory=list)
    duplicate_key: tuple = ()
    fatal: str = ""                  # message key khi cả file không chạy được

    def _count(self, status):
        return sum(1 for r in self.rows if r.status == status)

    @property
    def new_rows(self):
        return [r for r in self.rows if r.status == STATUS_NEW]

    @property
    def update_rows(self):
        return [r for r in self.rows if r.status == STATUS_UPDATE]

    @property
    def new_count(self):
        return self._count(STATUS_NEW)

    @property
    def update_count(self):
        return self._count(STATUS_UPDATE)

    @property
    def duplicate_count(self):
        return self._count(STATUS_DUPLICATE)

    @property
    def missing_count(self):
        return self._count(STATUS_MISSING)

    @property
    def skipped_count(self):
        return self.duplicate_count + self.missing_count

    @property
    def error_count(self):
        return self._count(STATUS_ERROR)

    @property
    def write_count(self):
        return self.new_count + self.update_count

    @property
    def problem_rows(self):
        return [r for r in self.rows if r.status == STATUS_ERROR]

    @property
    def can_apply(self):
        """Có lỗi thì KHÔNG cho ghi — sửa file rồi tải lại.

        Nửa vời (ghi dòng đúng, bỏ dòng sai) khiến lần nhập lại sau đó vừa trùng
        vừa thiếu, rất khó dò. Thà bắt sửa file một lần.
        """
        return (
            not self.fatal
            and not self.missing_columns
            and self.error_count == 0
            and self.write_count > 0
        )


def _clean_value(field, raw):
    value = raw
    if isinstance(value, str):
        value = value.strip()
    if isinstance(field, models.BooleanField):
        if isinstance(value, bool):
            return value
        text = str(value).strip().lower()
        if text in BOOL_TRUE:
            return True
        if text in BOOL_FALSE:
            return False if not field.null else (None if text == "" else False)
        raise ValidationError(f"'{raw}' không phải giá trị đúng/sai")
    if value in ("", None):
        if field.null:
            return None
        if field.has_default():
            return field.get_default()
        return "" if isinstance(field, (models.CharField, models.TextField)) else None
    return field.to_python(value)


def analyze(model, headers, rows, mode=DEFAULT_MODE):
    """Đối chiếu từng dòng với DB + validate, KHÔNG ghi gì."""
    mode = normalize_mode(mode)
    all_columns = columns_for(model)
    label = model_label(model)
    report = ImportReport(
        model_label=label,
        model_name=verbose_name(model),
        mode=mode,
        headers=list(headers),
        duplicate_key=natural_key_fields(model),
    )

    index_of = {}
    for position, header in enumerate(headers):
        index_of.setdefault(str(header).strip(), position)

    known_headers = {c.header for c in all_columns}
    report.unknown_columns = [h for h in headers if h and h not in known_headers]

    dup_key = report.duplicate_key
    if mode_writes_existing_rows(mode) and not dup_key:
        # Không có khoá thì không có cách nào biết dòng nào cần ghi đè.
        report.fatal = "admin.data.error.no_key_for_update"
        return report

    if mode == MODE_UPDATE:
        # Chỉ cần cột khoá. Cột nào có thì cập nhật, cột nào không có thì field
        # đó giữ nguyên giá trị cũ.
        columns = [c for c in all_columns if c.is_key or c.header in index_of]
        report.missing_columns = [c.header for c in all_columns if c.is_key and c.header not in index_of]
        report.ignored_columns = [c.header for c in all_columns if not c.is_key and c.header not in index_of]
    else:
        columns = all_columns
        report.missing_columns = [
            c.header for c in all_columns if c.required and c.header not in index_of
        ]
    if report.missing_columns:
        return report

    blank_keeps_old = BLANK_MEANS_UNCHANGED.get(label, ())
    hook = PRE_SAVE_HOOKS.get(label)
    seen_in_file = set()
    fk_cache = {}

    for offset, raw_row in enumerate(rows):
        number = offset + 2  # dòng 1 là header
        errors = []
        data = {}
        blank_fields = set()

        # 1. đổi từng ô sang giá trị Python / đối tượng FK
        for column in columns:
            position = index_of.get(column.header)
            raw = raw_row[position] if position is not None and position < len(raw_row) else ""
            if isinstance(raw, str):
                raw = raw.strip()
            field = model._meta.get_field(column.field_name)
            if column.is_fk:
                data.setdefault(f"__fk__{column.field_name}", {})[column.lookup or "pk"] = raw
                continue
            if str(raw) == "":
                blank_fields.add(column.field_name)
            try:
                data[column.field_name] = _clean_value(field, raw)
            except (ValidationError, ValueError, TypeError) as exc:
                errors.append(f"{column.header}: {_first_message(exc)}")

        # 2. dò khoá ngoại
        for key in [k for k in list(data) if k.startswith("__fk__")]:
            field_name = key[len("__fk__"):]
            lookups = data.pop(key)
            field = model._meta.get_field(field_name)
            filters = {k: v for k, v in lookups.items() if str(v).strip() != ""}
            if not filters:
                if field.null or field.blank:
                    data[field_name] = None
                else:
                    errors.append(f"{field_name}: thiếu tham chiếu bắt buộc")
                continue
            cache_key = (field.related_model, tuple(sorted(filters.items())))
            if cache_key in fk_cache:
                related = fk_cache[cache_key]
            else:
                try:
                    related = field.related_model._default_manager.get(**filters)
                except field.related_model.DoesNotExist:
                    related = None
                except Exception as exc:  # nhiều bản ghi khớp / lookup sai kiểu
                    related = None
                    errors.append(f"{field_name}: {_first_message(exc)}")
                fk_cache[cache_key] = related
            if related is None:
                shown = ", ".join(f"{k}={v}" for k, v in filters.items())
                errors.append(f"{field_name}: không tìm thấy bản ghi khớp ({shown})")
            else:
                data[field_name] = related

        if errors:
            report.rows.append(RowResult(number=number, status=STATUS_ERROR, errors=errors))
            continue

        # 3. dòng này đã có trong file / trong DB chưa?
        existing = None
        if dup_key:
            try:
                signature = tuple(_signature_value(data.get(name)) for name in dup_key)
            except Exception:  # pragma: no cover
                signature = None
            if signature is not None:
                if signature in seen_in_file:
                    report.rows.append(
                        RowResult(
                            number=number,
                            status=STATUS_DUPLICATE,
                            display=_display(dup_key, data),
                            errors=["trùng với một dòng phía trên trong cùng file"],
                        )
                    )
                    continue
                seen_in_file.add(signature)
                filters = {name: data.get(name) for name in dup_key}
                existing = model._default_manager.filter(**filters).first()

        if existing is not None and not mode_writes_existing_rows(mode):
            report.rows.append(
                RowResult(
                    number=number,
                    status=STATUS_DUPLICATE,
                    display=_display(dup_key, data),
                    errors=["đã có trong hệ thống"],
                )
            )
            continue
        if existing is None and not mode_writes_new_rows(mode):
            report.rows.append(
                RowResult(
                    number=number,
                    status=STATUS_MISSING,
                    display=_display(dup_key, data),
                    errors=["chưa có trong hệ thống nên không có gì để cập nhật"],
                )
            )
            continue

        # 4. dựng bản ghi rồi validate bằng chính model
        if existing is None:
            instance = model(**data)
        else:
            instance = existing
            for field_name, value in data.items():
                if field_name in blank_fields and field_name in blank_keeps_old:
                    continue
                setattr(instance, field_name, value)
        if hook:
            hook(instance, data, is_new=existing is None)
        exclude = [f.name for f in model._meta.fields if f.name in AUDIT_FIELDS or f.primary_key]
        exclude += EXTRA_VALIDATION_EXCLUDES.get(label, [])
        if existing is not None:
            # Cột không có trong file thì field đó không được validate lại —
            # dữ liệu cũ trong DB có thể đã được nhập từ thời quy tắc khác.
            touched = set(data)
            exclude += [
                f.name
                for f in model._meta.fields
                if f.editable and not f.primary_key and f.name not in touched
            ]
        try:
            instance.full_clean(exclude=list(dict.fromkeys(exclude)))
        except ValidationError as exc:
            report.rows.append(
                RowResult(number=number, status=STATUS_ERROR, errors=_flatten(exc))
            )
            continue

        report.rows.append(
            RowResult(
                number=number,
                status=STATUS_NEW if existing is None else STATUS_UPDATE,
                display=_display(dup_key, data) or str(instance),
                data=data,
                pk=None if existing is None else existing.pk,
            )
        )
    return report


def _signature_value(value):
    if isinstance(value, models.Model):
        return value.pk
    if isinstance(value, str):
        return value.strip()
    return value


def _display(dup_key, data):
    if not dup_key:
        return ""
    parts = []
    for name in dup_key:
        value = data.get(name)
        parts.append(str(getattr(value, "pk", value)))
    return " / ".join(parts)


def _first_message(exc):
    if isinstance(exc, ValidationError):
        messages = _flatten(exc)
        return messages[0] if messages else str(exc)
    return str(exc)


def _flatten(exc):
    out = []
    if hasattr(exc, "message_dict"):
        for field_name, messages in exc.message_dict.items():
            for m in messages:
                out.append(m if field_name == "__all__" else f"{field_name}: {m}")
    else:
        out.extend(exc.messages if hasattr(exc, "messages") else [str(exc)])
    return out


def apply_report(report, model):
    """Ghi các dòng 'mới' và 'cập nhật'. GỌI TRONG transaction.atomic() ở tầng view.

    Dùng create()/save() từng dòng chứ không bulk_create: AuditableModel.save()
    mới điền created_by/updated_by từ CurrentUserMiddleware, bulk_create bỏ qua save().
    """
    label = model_label(model)
    hook = PRE_SAVE_HOOKS.get(label)
    blank_keeps_old = BLANK_MEANS_UNCHANGED.get(label, ())
    created = 0
    updated = 0
    for row in report.rows:
        if row.status == STATUS_NEW:
            instance = model(**row.data)
            if hook:
                hook(instance, row.data, is_new=True)
            instance.save()
            created += 1
        elif row.status == STATUS_UPDATE:
            instance = model._default_manager.get(pk=row.pk)
            for field_name, value in row.data.items():
                if field_name in blank_keeps_old and not value:
                    continue
                setattr(instance, field_name, value)
            if hook:
                hook(instance, row.data, is_new=False)
            instance.save()
            updated += 1
    return {"created": created, "updated": updated}


# ---------------------------------------------------------------------------
# Hook riêng theo bảng
# ---------------------------------------------------------------------------

def _prepare_user(instance, data, is_new=True):
    """Không bao giờ ghi mật khẩu thô xuống DB.

    Ô password trong file: để trống + dòng MỚI -> tài khoản không đăng nhập được
    bằng mật khẩu (đúng ý khi nhập danh sách người học rồi bắt họ đặt lại). Để
    trống + dòng CẬP NHẬT -> giữ nguyên mật khẩu cũ, tuyệt đối không khoá tài
    khoản người ta chỉ vì file thiếu một ô. Có giá trị và chưa băm -> băm ngay
    tại đây. Đã là chuỗi băm (copy từ bảng khác) -> giữ nguyên.
    """
    from django.contrib.auth.hashers import identify_hasher, make_password

    raw = (data.get("password") or "").strip()
    if not raw:
        if is_new:
            instance.set_unusable_password()
        return
    try:
        identify_hasher(raw)
    except Exception:
        instance.password = make_password(raw)


PRE_SAVE_HOOKS = {
    "accounts.user": _prepare_user,
}

# password của User được hook xử lý riêng nên bỏ khỏi full_clean (nếu không,
# ô trống sẽ báo "không được để trống" dù ta cố ý đặt mật khẩu không dùng được).
EXTRA_VALIDATION_EXCLUDES = {
    "accounts.user": ["password"],
}


# ===========================================================================
# MẪU GỘP — "Từ vựng đầy đủ"
# ===========================================================================
# Một từ vựng nằm trên 3 bảng: Vocabulary (mặt chữ/cách đọc/nghĩa),
# VocabularyTopic (chủ đề) và ExampleSentence (câu ví dụ). Nhập theo từng bảng
# nghĩa là tải lên 3 file đúng thứ tự — đúng nhưng không ai muốn làm khi thêm
# 200 từ mới. Bộ cột dưới đây gộp cả 3 vào MỘT dòng.
#
# Quan hệ con là CỘNG THÊM, KHÔNG THAY THẾ: chủ đề và câu ví dụ trong file được
# thêm vào nếu chưa có, và không bao giờ gỡ bớt thứ đang có. Chọn như vậy vì gỡ
# liên kết là thao tác không hoàn tác được mà một ô bỏ trống nhầm cũng kích
# hoạt. Muốn bỏ một chủ đề khỏi một từ thì xoá dòng đó ở Django admin
# (/admin/vocabulary/vocabularytopic/).

FULL_VOCAB_LABEL = "vocabulary.vocabulary_full"
FULL_VOCAB_EXAMPLE_SLOTS = 3
FULL_VOCAB_TOPIC_COLUMN = "topics"
# Slug không chứa dấu ; , hay xuống dòng nên tách bằng cả ba đều an toàn.
_TOPIC_SPLIT = re.compile(r"[;,\n]+")


def _full_vocab_models():
    from apps.vocabulary.models import ExampleSentence, Topic, Vocabulary, VocabularyTopic

    return Vocabulary, Topic, VocabularyTopic, ExampleSentence


def full_vocab_columns():
    Vocabulary, Topic, _, _ = _full_vocab_models()
    columns = []
    for column in columns_for(Vocabulary):
        columns.append(column)
    columns.append(
        Column(
            header=FULL_VOCAB_TOPIC_COLUMN,
            field_name=FULL_VOCAB_TOPIC_COLUMN,
            lookup="slug",
            is_fk=True,
            required=False,
            is_key=False,
            hint=(
                "Danh sách slug chủ đề, ngăn nhau bằng dấu ; — vd "
                "'hop-hanh;dien-thoai'. Chủ đề phải có sẵn (nhập bảng "
                "vocabulary.topic trước). Chỉ THÊM liên kết còn thiếu, không gỡ "
                "liên kết đang có."
            ),
        )
    )
    for slot in range(1, FULL_VOCAB_EXAMPLE_SLOTS + 1):
        columns.append(
            Column(
                header=f"example{slot}_jp",
                field_name=f"example{slot}_jp",
                lookup=None,
                is_fk=False,
                required=False,
                is_key=False,
                hint=f"Câu ví dụ tiếng Nhật thứ {slot}. Điền thì phải điền cả example{slot}_vi.",
            )
        )
        columns.append(
            Column(
                header=f"example{slot}_vi",
                field_name=f"example{slot}_vi",
                lookup=None,
                is_fk=False,
                required=False,
                is_key=False,
                hint=f"Nghĩa tiếng Việt của câu ví dụ thứ {slot}.",
            )
        )
    return columns


def full_vocab_headers():
    return [c.header for c in full_vocab_columns()]


def full_vocab_guide_rows():
    Vocabulary, _, VocabularyTopic, ExampleSentence = _full_vocab_models()
    out = [["Cột", "Bắt buộc", "Ghi chú"]]
    for column in full_vocab_columns():
        note = column.hint
        if column.is_key:
            note = f"[KHOÁ] {note}"
        out.append([column.header, "x" if column.required else "", note])
    out.extend(_guide_mode_rows())
    out.extend(
        [
            [],
            ["Khoá chống trùng", "", "word, reading"],
            [
                "Ghi vào 3 bảng",
                "",
                f"{model_label(Vocabulary)} + {model_label(VocabularyTopic)} + "
                f"{model_label(ExampleSentence)}, tất cả trong một transaction.",
            ],
            [
                "Chủ đề & câu ví dụ",
                "",
                "CHỈ THÊM, không gỡ. Muốn bỏ bớt thì xoá ở trang quản trị Django.",
            ],
        ]
    )
    return out


def full_vocab_export_rows():
    Vocabulary, _, _, _ = _full_vocab_models()
    base_columns = columns_for(Vocabulary)
    rows = []
    queryset = (
        Vocabulary._default_manager.all()
        .prefetch_related("topic_links__topic", "examples")
        .order_by("word", "reading")
    )
    for obj in queryset:
        row = [getattr(obj, c.field_name, "") or "" for c in base_columns]
        slugs = sorted(link.topic.slug for link in obj.topic_links.all() if link.topic_id)
        row.append(";".join(slugs))
        examples = list(obj.examples.all().order_by("pk"))[:FULL_VOCAB_EXAMPLE_SLOTS]
        for slot in range(FULL_VOCAB_EXAMPLE_SLOTS):
            if slot < len(examples):
                row.append(examples[slot].sentence_jp)
                row.append(examples[slot].sentence_vi)
            else:
                row.extend(["", ""])
        rows.append(row)
    return rows


def analyze_full_vocab(headers, rows, mode=DEFAULT_MODE):
    """Như analyze() nhưng cho mẫu gộp. Dựng sẵn cả phần chủ đề và câu ví dụ."""
    mode = normalize_mode(mode)
    Vocabulary, Topic, _, _ = _full_vocab_models()
    all_columns = full_vocab_columns()
    base_columns = columns_for(Vocabulary)
    report = ImportReport(
        model_label=FULL_VOCAB_LABEL,
        model_name="Từ vựng đầy đủ",
        mode=mode,
        headers=list(headers),
        duplicate_key=("word", "reading"),
    )

    index_of = {}
    for position, header in enumerate(headers):
        index_of.setdefault(str(header).strip(), position)

    known_headers = {c.header for c in all_columns}
    report.unknown_columns = [h for h in headers if h and h not in known_headers]

    if mode == MODE_UPDATE:
        columns = [c for c in all_columns if c.is_key or c.header in index_of]
        report.missing_columns = [c.header for c in all_columns if c.is_key and c.header not in index_of]
        report.ignored_columns = [c.header for c in all_columns if not c.is_key and c.header not in index_of]
    else:
        columns = all_columns
        report.missing_columns = [c.header for c in all_columns if c.required and c.header not in index_of]
    if report.missing_columns:
        return report

    active = {c.header for c in columns}
    topic_cache = {}
    seen_in_file = set()

    def cell(raw_row, header):
        position = index_of.get(header)
        if position is None or position >= len(raw_row):
            return ""
        value = raw_row[position]
        return value.strip() if isinstance(value, str) else ("" if value is None else str(value).strip())

    for offset, raw_row in enumerate(rows):
        number = offset + 2
        errors = []
        base_data = {}

        for column in base_columns:
            if column.header not in active:
                continue
            field = Vocabulary._meta.get_field(column.field_name)
            try:
                base_data[column.field_name] = _clean_value(field, cell(raw_row, column.header))
            except (ValidationError, ValueError, TypeError) as exc:
                errors.append(f"{column.header}: {_first_message(exc)}")

        # chủ đề
        topics = []
        topics_present = FULL_VOCAB_TOPIC_COLUMN in active
        if topics_present:
            for slug in _TOPIC_SPLIT.split(cell(raw_row, FULL_VOCAB_TOPIC_COLUMN)):
                slug = slug.strip()
                if not slug:
                    continue
                if slug not in topic_cache:
                    topic_cache[slug] = Topic._default_manager.filter(slug=slug).first()
                topic = topic_cache[slug]
                if topic is None:
                    errors.append(f"{FULL_VOCAB_TOPIC_COLUMN}: không có chủ đề slug='{slug}'")
                elif topic not in topics:
                    topics.append(topic)

        # câu ví dụ
        examples = []
        for slot in range(1, FULL_VOCAB_EXAMPLE_SLOTS + 1):
            jp_header, vi_header = f"example{slot}_jp", f"example{slot}_vi"
            if jp_header not in active and vi_header not in active:
                continue
            jp = cell(raw_row, jp_header)
            vi = cell(raw_row, vi_header)
            if not jp and not vi:
                continue
            if not jp or not vi:
                errors.append(f"example{slot}: phải điền cả câu tiếng Nhật và nghĩa tiếng Việt")
                continue
            if len(jp) > 255 or len(vi) > 255:
                errors.append(f"example{slot}: câu ví dụ tối đa 255 ký tự")
                continue
            examples.append((jp, vi))

        if errors:
            report.rows.append(RowResult(number=number, status=STATUS_ERROR, errors=errors))
            continue

        word = base_data.get("word", "")
        reading = base_data.get("reading", "")
        signature = (str(word).strip(), str(reading).strip())
        if signature in seen_in_file:
            report.rows.append(
                RowResult(
                    number=number,
                    status=STATUS_DUPLICATE,
                    display=f"{word} / {reading}",
                    errors=["trùng với một dòng phía trên trong cùng file"],
                )
            )
            continue
        seen_in_file.add(signature)

        existing = Vocabulary._default_manager.filter(word=word, reading=reading).first()
        if existing is not None and not mode_writes_existing_rows(mode):
            report.rows.append(
                RowResult(
                    number=number,
                    status=STATUS_DUPLICATE,
                    display=f"{word} / {reading}",
                    errors=["đã có trong hệ thống"],
                )
            )
            continue
        if existing is None and not mode_writes_new_rows(mode):
            report.rows.append(
                RowResult(
                    number=number,
                    status=STATUS_MISSING,
                    display=f"{word} / {reading}",
                    errors=["chưa có trong hệ thống nên không có gì để cập nhật"],
                )
            )
            continue

        if existing is None:
            instance = Vocabulary(**base_data)
        else:
            instance = existing
            for field_name, value in base_data.items():
                setattr(instance, field_name, value)
        exclude = [f.name for f in Vocabulary._meta.fields if f.name in AUDIT_FIELDS or f.primary_key]
        if existing is not None:
            touched = set(base_data)
            exclude += [
                f.name
                for f in Vocabulary._meta.fields
                if f.editable and not f.primary_key and f.name not in touched
            ]
        try:
            instance.full_clean(exclude=list(dict.fromkeys(exclude)))
        except ValidationError as exc:
            report.rows.append(RowResult(number=number, status=STATUS_ERROR, errors=_flatten(exc)))
            continue

        report.rows.append(
            RowResult(
                number=number,
                status=STATUS_NEW if existing is None else STATUS_UPDATE,
                display=f"{word} / {reading}",
                data={"base": base_data, "topics": topics, "examples": examples},
                pk=None if existing is None else existing.pk,
            )
        )
    return report


def apply_full_vocab(report):
    """Ghi mẫu gộp. GỌI TRONG transaction.atomic().

    `objects.create()` từng dòng, không bulk_create — xem apply_report().
    """
    Vocabulary, _, VocabularyTopic, ExampleSentence = _full_vocab_models()
    created = updated = topics_linked = examples_added = 0

    for row in report.rows:
        if row.status not in (STATUS_NEW, STATUS_UPDATE):
            continue
        base = row.data["base"]
        if row.status == STATUS_NEW:
            vocabulary = Vocabulary(**base)
            vocabulary.save()
            created += 1
        else:
            vocabulary = Vocabulary._default_manager.get(pk=row.pk)
            for field_name, value in base.items():
                setattr(vocabulary, field_name, value)
            vocabulary.save()
            updated += 1

        for topic in row.data["topics"]:
            if not VocabularyTopic._default_manager.filter(
                vocabulary=vocabulary, topic=topic
            ).exists():
                VocabularyTopic._default_manager.create(vocabulary=vocabulary, topic=topic)
                topics_linked += 1

        for sentence_jp, sentence_vi in row.data["examples"]:
            example = ExampleSentence._default_manager.filter(
                vocabulary=vocabulary, sentence_jp=sentence_jp
            ).first()
            if example is None:
                ExampleSentence._default_manager.create(
                    vocabulary=vocabulary, sentence_jp=sentence_jp, sentence_vi=sentence_vi
                )
                examples_added += 1
            elif example.sentence_vi != sentence_vi and mode_writes_existing_rows(report.mode):
                example.sentence_vi = sentence_vi
                example.save()

    return {
        "created": created,
        "updated": updated,
        "topics_linked": topics_linked,
        "examples_added": examples_added,
    }


# ---------------------------------------------------------------------------
# Dataset — lớp bọc để view không phải phân biệt "bảng thường" và "mẫu gộp"
# ---------------------------------------------------------------------------

class ModelDataset:
    """Một bảng thật trong DB."""

    is_composite = False

    def __init__(self, model):
        self.model = model
        self.label = model_label(model)
        self.name = verbose_name(model)
        self.app = model._meta.app_label

    @property
    def perm_models(self):
        return [self.model]

    def count(self):
        return self.model._default_manager.count()

    def columns(self):
        return columns_for(self.model)

    def headers(self):
        return headers_for(self.model)

    def guide_rows(self):
        return guide_rows(self.model)

    def export_rows(self):
        return export_rows(self.model)

    def natural_key(self):
        return natural_key_fields(self.model)

    def supports_update(self):
        return bool(self.natural_key())

    def m2m_notes(self):
        return skipped_m2m_names(self.model)

    def analyze(self, headers, rows, mode=DEFAULT_MODE):
        return analyze(self.model, headers, rows, mode=mode)

    def apply(self, report):
        return apply_report(report, self.model)


class FullVocabularyDataset:
    """Mẫu gộp: Vocabulary + VocabularyTopic + ExampleSentence trên một dòng."""

    is_composite = True
    label = FULL_VOCAB_LABEL
    name = "Từ vựng đầy đủ (kèm chủ đề & câu ví dụ)"
    app = "vocabulary"

    @property
    def perm_models(self):
        Vocabulary, _, VocabularyTopic, ExampleSentence = _full_vocab_models()
        return [Vocabulary, VocabularyTopic, ExampleSentence]

    def count(self):
        Vocabulary, _, _, _ = _full_vocab_models()
        return Vocabulary._default_manager.count()

    def columns(self):
        return full_vocab_columns()

    def headers(self):
        return full_vocab_headers()

    def guide_rows(self):
        return full_vocab_guide_rows()

    def export_rows(self):
        return full_vocab_export_rows()

    def natural_key(self):
        return ("word", "reading")

    def supports_update(self):
        return True

    def m2m_notes(self):
        return []

    def analyze(self, headers, rows, mode=DEFAULT_MODE):
        return analyze_full_vocab(headers, rows, mode=mode)

    def apply(self, report):
        return apply_full_vocab(report)


def datasets():
    """Mẫu gộp đứng đầu, rồi tới 18 bảng theo thứ tự phụ thuộc."""
    return [FullVocabularyDataset()] + [ModelDataset(m) for m in importable_models()]


def get_dataset(label):
    label = (label or "").strip().lower()
    if label == FULL_VOCAB_LABEL:
        return FullVocabularyDataset()
    model = get_importable_model(label)
    return ModelDataset(model) if model is not None else None
